#!/usr/bin/env python3
"""Convert the raw downloads in data/raw/ into small, plain TSV files in data/extract/.

The extract files are what build_bands.py reads, so the build itself needs only the
standard library. This step needs `openpyxl` (xlsx) and `pdftotext` (poppler-utils).

    python3 scripts/extract_raw.py

Outputs (all UTF-8, tab-separated, header row):
  data/extract/awl_families.tsv      sublist, headword, members
  data/extract/subtlex_zipf.tsv      word, zipf, dom_pos, all_pos
  data/extract/prevalence.tsv        word, pknown, prevalence, nobs
  data/extract/oxford_cefr.tsv       word, pos, cefr, cefr_all, list
  data/extract/nation_pdf_headwords.tsv  band, headword   (official PDFs, for verification)
  data/extract/oewn_senses.tsv       lemma, pos, sense_no, synset, definition, example
  data/extract/oewn_relations.tsv    lemma, pos, sense_no, relation, target
  data/extract/oewn_forms.tsv        form, lemma   (inflected forms OEWN lists, e.g. media -> medium)
    (Open English WordNet, only lemmas that are a family headword or member in vocab/index.csv)
"""
from __future__ import annotations

import csv
import gzip
import html
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "extract"


def write_tsv(path: Path, header: list[str], rows) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter="\t", lineterminator="\n")
        w.writerow(header)
        n = 0
        for r in rows:
            w.writerow(r)
            n += 1
    print(f"  wrote {path.relative_to(ROOT)}  ({n} rows)")
    return n


# ---------------------------------------------------------------- AWL
# The official sublist pages omit `confirm` from sublist 7 (59 headwords instead of 60).
AWL_FIXES = {
    7: [("confirm", "confirmation|confirmations|confirmed|confirming|confirms")],
}


def extract_awl():
    rows = []
    for i in range(1, 11):
        s = (RAW / "awl" / f"sublist{i:02d}.html").read_text(encoding="utf-8")
        body = s[s.find("<h2>The Academic Word List</h2>"):]
        body = body[: body.find("</main>")]
        for m in re.finditer(r"<p>([^<]+)</p>(?:\s*<ul>(.*?)</ul>)?", body, re.S):
            hw = html.unescape(m.group(1)).strip().lower().replace(" ", "-")
            members = [html.unescape(x).strip().lower() for x in re.findall(r"<li>([^<]+)</li>", m.group(2) or "")]
            rows.append((i, hw, "|".join(members)))
        for hw, members in AWL_FIXES.get(i, []):
            rows.append((i, hw, members))
    rows.sort(key=lambda r: (r[0], r[1]))
    write_tsv(OUT / "awl_families.tsv", ["sublist", "headword", "members"], rows)


# ---------------------------------------------------------------- SUBTLEX-US
def extract_subtlex():
    import openpyxl

    wb = openpyxl.load_workbook(RAW / "subtlex" / "SUBTLEX-US frequency list with PoS and Zipf information.xlsx", read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    iw, iz = header.index("Word"), header.index("Zipf-value")
    idp, iap = header.index("Dom_PoS_SUBTLEX"), header.index("All_PoS_SUBTLEX")

    def rows():
        for r in it:
            if r[iw] is None:
                continue
            yield (str(r[iw]).strip().lower(), f"{float(r[iz]):.3f}", r[idp] or "", (r[iap] or "").replace(".", "|"))

    write_tsv(OUT / "subtlex_zipf.tsv", ["word", "zipf", "dom_pos", "all_pos"], rows())


# ---------------------------------------------------------------- prevalence
def extract_prevalence():
    import openpyxl

    wb = openpyxl.load_workbook(RAW / "prevalence" / "English_Word_Prevalences.xlsx", read_only=True)
    ws = wb["Prevalence"]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    iw, ip, ipr, inobs = header.index("Word"), header.index("Pknown"), header.index("Prevalence"), header.index("Nobs")

    def rows():
        for r in it:
            if r[iw] is None:
                continue
            yield (str(r[iw]).strip().lower(), f"{float(r[ip]):.4f}", f"{float(r[ipr]):.3f}", r[inobs])

    write_tsv(OUT / "prevalence.tsv", ["word", "pknown", "prevalence", "nobs"], rows())


# ---------------------------------------------------------------- Oxford 3000/5000
POS = (r"(?:n\.|v\.|adj\.|adv\.|prep\.|det\.|pron\.|conj\.|exclam\.|number|modal v\.|auxiliary v\.|"
       r"indefinite article|definite article|infinitive marker|adj\./adv\.|det\./pron\.|conj\./prep\.|prefix|suffix|abbr\.)")
LEV = r"\b([ABC][12])\b"
SKIP = ("The Oxford", "most important words", "expanded core", "additional 2000", "3000, it includes", "Oxford University Press")


def pdftotext(pdf: Path) -> str:
    return subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True).stdout


def parse_oxford(text: str):
    """Yield (word, pos, min_level, all_levels) for each entry cell in the layout text."""
    cells: list[str] = []
    for line in text.splitlines():
        if not line.strip() or any(k in line for k in SKIP) or re.fullmatch(r"\s*\d+\s*/\s*\d+\s*", line):
            continue
        cells += [c.strip() for c in re.split(r"\s{2,}", line.strip()) if c.strip()]
    # Entries can wrap onto the next line of the same column: a cell ending in ',' continues
    # in the next cell that starts with a PoS token (e.g. "double adj., det., pron., v. A2," +
    # "adv. B1"). Cells from the other columns come in between, so keep the wrapped cell aside.
    merged: list[str] = []
    pending = None
    for cell in cells:
        if pending is not None and re.match(POS, cell):
            merged.append(pending + " " + cell)
            pending = None
        elif cell.endswith(","):
            if pending is not None:
                merged.append(pending)
            pending = cell
        else:
            merged.append(cell)
    if pending is not None:
        merged.append(pending)
    for cell in merged:
        levels = re.findall(LEV, cell)
        m = re.search(r"\s(?=" + POS + r")|\s(?=" + LEV + ")", cell)
        if not levels or not m:
            continue
        word = cell[: m.start()].strip()
        word = re.sub(r"\s*\([^)]*\)", "", word)   # drop sense glosses: kind (type)
        word = re.sub(r"(?<=[a-z])\d$", "", word)  # drop homograph numbers: wind1
        pos = [p.rstrip(".") for p in re.findall(POS, cell)]
        yield word.lower(), "|".join(dict.fromkeys(pos)), min(levels), "|".join(levels)


def extract_oxford():
    merged: dict[str, dict] = {}
    for name, tag in (("The_Oxford_3000", "oxford3000"), ("The_Oxford_5000", "oxford5000")):
        for word, pos, lev, levs in parse_oxford(pdftotext(RAW / "oxford" / f"{name}.pdf")):
            e = merged.setdefault(word, {"pos": [], "levels": [], "list": tag})
            e["pos"] += [p for p in pos.split("|") if p and p not in e["pos"]]
            e["levels"] += levs.split("|")
    rows = sorted((w, "|".join(e["pos"]), min(e["levels"]), "|".join(sorted(set(e["levels"]))), e["list"]) for w, e in merged.items())
    write_tsv(OUT / "oxford_cefr.tsv", ["word", "pos", "cefr", "cefr_all", "list"], rows)


# ---------------------------------------------------------------- Nation official headword PDFs
def extract_nation_pdf():
    rows = []
    for band, n in enumerate(("first", "second", "third", "fourth", "fifth", "sixth"), start=1):
        text = pdftotext(RAW / "nation" / "official-pdf" / f"headwords-{n}-thousand.pdf")
        seen = set()
        for line in text.splitlines():
            if "Headwords of the" in line:
                continue
            for tok in re.split(r"\s+", line.strip()):
                if re.fullmatch(r"[A-Z][A-Z'-]*", tok) and tok not in seen:
                    seen.add(tok)
                    rows.append((band, tok.lower()))
    write_tsv(OUT / "nation_pdf_headwords.tsv", ["band", "headword"], rows)


# ---------------------------------------------------------------- Open English WordNet
OEWN_POS = {"n": "n", "v": "v", "a": "adj", "s": "adj", "r": "adv"}


def extract_oewn():
    """Slice OEWN down to our families: one row per sense, one row per synonym/antonym/similar link.

    Sense order inside a lemma+pos is the document order of <Sense> elements (WordNet sense
    order, most frequent first). Relation targets are written as lemmas, so links that point
    outside the 6,000 families are still visible to the audit.
    """
    wanted = set()
    with (ROOT / "vocab" / "index.csv").open(encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            wanted.add(r["family"])
            wanted.update(m for m in r["members"].split("|") if m)
    wanted |= {w.capitalize() for w in wanted}       # Friday, January, Christmas

    entries: dict[str, tuple[str, str, list[str]]] = {}      # entry id -> (lemma, pos, [sense id])
    forms: set[tuple[str, str]] = set()                        # (inflected form, lemma) for wanted forms
    senses: dict[str, tuple[str, str, int, str, list[str]]] = {}  # sense id -> (lemma, pos, sense_no, synset, [antonym sense id])
    synsets: dict[str, tuple[str, str, list[str], list[str]]] = {}  # synset id -> (definition, example, [member entry id], [similar synset id])

    with gzip.open(RAW / "oewn" / "english-wordnet-2025.xml.gz", "rb") as fh:
        for _, el in ET.iterparse(fh):
            if el.tag == "LexicalEntry":
                lemma_el = el.find("Lemma")
                lemma, pos = lemma_el.get("writtenForm"), lemma_el.get("partOfSpeech")
                ids = []
                for i, s in enumerate(el.findall("Sense"), start=1):
                    ants = [r.get("target") for r in s.findall("SenseRelation") if r.get("relType") == "antonym"]
                    senses[s.get("id")] = (lemma, pos, i, s.get("synset"), ants)
                    ids.append(s.get("id"))
                entries[el.get("id")] = (lemma, pos, ids)
                forms.update((f.get("writtenForm"), lemma) for f in el.findall("Form") if f.get("writtenForm") in wanted)
                el.clear()
            elif el.tag == "Synset":
                d = el.find("Definition")
                ex = el.find("Example")
                sim = [r.get("target") for r in el.findall("SynsetRelation") if r.get("relType") == "similar"]
                synsets[el.get("id")] = ((d.text or "").strip() if d is not None else "",
                                         (ex.text or "").strip() if ex is not None else "",
                                         el.get("members", "").split(), sim)
                el.clear()

    sense_rows, rel_rows = [], []
    for sid, (lemma, pos, n, syn, ants) in senses.items():
        if lemma not in wanted or pos not in OEWN_POS:
            continue
        definition, example, members, similar = synsets[syn]
        tag = OEWN_POS[pos]
        sense_rows.append((lemma, tag, n, syn, definition, example))
        for eid in members:
            other = entries[eid][0]
            if other != lemma:
                rel_rows.append((lemma, tag, n, "synonym", other))
        for tid in ants:
            rel_rows.append((lemma, tag, n, "antonym", senses[tid][0]))
        for ts in similar:
            for eid in synsets[ts][2]:
                rel_rows.append((lemma, tag, n, "similar", entries[eid][0]))
    sense_rows.sort(key=lambda r: (r[0], r[1], r[2]))
    rel_rows.sort(key=lambda r: (r[0], r[1], r[2], r[3], r[4]))
    write_tsv(OUT / "oewn_senses.tsv", ["lemma", "pos", "sense_no", "synset", "definition", "example"], sense_rows)
    write_tsv(OUT / "oewn_relations.tsv", ["lemma", "pos", "sense_no", "relation", "target"], rel_rows)
    write_tsv(OUT / "oewn_forms.tsv", ["form", "lemma"], sorted(forms))


if __name__ == "__main__":
    steps = {"awl": extract_awl, "subtlex": extract_subtlex, "prevalence": extract_prevalence,
             "oxford": extract_oxford, "nation_pdf": extract_nation_pdf, "oewn": extract_oewn}
    wanted = sys.argv[1:] or list(steps)
    for name in wanted:
        print(f"== {name}")
        steps[name]()
